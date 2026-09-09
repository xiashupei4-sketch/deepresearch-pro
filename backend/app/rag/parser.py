"""Document parsing: PDF / DOCX / TXT / Markdown → plain text."""

from __future__ import annotations

import io
from pathlib import Path


def parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())
    return "\n\n".join(parts)


def parse_docx(data: bytes) -> str:
    import docx  # python-docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def parse_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")


def parse_document(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(data)
    if suffix == ".docx":
        return parse_docx(data)
    if suffix in (".txt", ".md", ".markdown"):
        return parse_txt(data)
    raise ValueError(f"Unsupported document type: {suffix}")


SUPPORTED_TYPES = {".pdf", ".docx", ".txt", ".md", ".markdown"}
